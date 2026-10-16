"""Local SSH/Telnet fixtures only; the shipped collector has no Python dependency."""
from __future__ import annotations

import csv
import json
import re
import socket
import subprocess
import threading
import time
import unittest
import zipfile
from pathlib import Path

import paramiko
from openpyxl import Workbook
from telnet_fixture import TelnetDeviceServer

ROOT = Path(__file__).resolve().parents[1]
EXE = ROOT / "build" / "netcollect-cli.exe"
ARTIFACTS = ROOT / "test-output"
KEY = paramiko.RSAKey.generate(2048)
RUN_GUI = False
COMMANDS = ["display version", "display interface | include up", "display big", "display slow"]


class Auth(paramiko.ServerInterface):
    def __init__(self, owner):
        self.owner = owner
        self.shell = threading.Event()

    def get_allowed_auths(self, username):
        return "keyboard-interactive" if self.owner.keyboard else "password"

    def check_auth_password(self, username, password):
        if not self.owner.keyboard and username == "admin" and password == " secret ":
            return paramiko.AUTH_SUCCESSFUL
        return paramiko.AUTH_FAILED

    def check_auth_interactive(self, username, submethods):
        return paramiko.InteractiveQuery("Login", "", ("Password: ", False))

    def check_auth_interactive_response(self, responses):
        return paramiko.AUTH_SUCCESSFUL if responses == [" secret "] else paramiko.AUTH_FAILED

    def check_channel_request(self, kind, chanid):
        return paramiko.OPEN_SUCCEEDED if kind == "session" else paramiko.OPEN_FAILED_ADMINISTRATIVELY_PROHIBITED

    def check_channel_pty_request(self, *args):
        return True

    def check_channel_shell_request(self, channel):
        self.shell.set()
        return True


class DeviceServer:
    def __init__(self, name, style="angle", encoding="utf-8", keyboard=False, key=KEY, address="127.0.0.1"):
        self.name, self.style, self.encoding, self.keyboard, self.key = name, style, encoding, keyboard, key
        self.commands = []
        self.timestamps = []
        self.errors = []
        self.stop = threading.Event()
        self.address = address
        self.sock = socket.socket(socket.AF_INET6 if ":" in address else socket.AF_INET)
        self.sock.bind((address, 0))
        self.sock.listen(8)
        self.sock.settimeout(0.2)
        self.port = self.sock.getsockname()[1]
        self.connections = []
        self.threads = []
        self.thread = threading.Thread(target=self.accept, daemon=True)
        self.thread.start()

    @property
    def prompt(self):
        if self.style == "angle":
            return f"<{self.name}>"
        if self.style == "square":
            return f"[{self.name}]"
        return f"{self.name}#"

    def accept(self):
        while not self.stop.is_set():
            try:
                client, _ = self.sock.accept()
            except socket.timeout:
                continue
            except OSError:
                break
            t = threading.Thread(target=self.serve, args=(client,), daemon=True)
            self.threads.append(t)
            t.start()

    def serve(self, client):
        transport = paramiko.Transport(client)
        self.connections.append(transport)
        try:
            transport.add_server_key(self.key)
            auth = Auth(self)
            transport.start_server(server=auth)
            channel = transport.accept(8)
            if channel is None or not auth.shell.wait(5):
                return
            channel.settimeout(8)

            def send(text):
                channel.sendall(text.encode(self.encoding))

            send("\x1b[32mWelcome\x1b[0m\r\n")
            time.sleep(0.03)
            send(self.prompt)
            pending = b""
            while transport.is_active() and not self.stop.is_set():
                chunk = channel.recv(1024)
                if not chunk:
                    break
                pending += chunk
                while b"\n" in pending:
                    command, pending = pending.split(b"\n", 1)
                    command = command.decode("utf-8").strip()
                    self.commands.append(command)
                    self.timestamps.append(time.perf_counter())
                    send(command + "\r\n")
                    if command == "display version":
                        message = "設備運行正常，接口狀態正常\r\n" if self.encoding == "big5" else "设备运行正常，接口状态正常\r\n"
                        data = message.encode(self.encoding)
                        channel.sendall(data[:3])
                        time.sleep(0.02)
                        channel.sendall(data[3:])
                    elif command == "display interface | include up":
                        send("GigabitEthernet0/0 up\r\n--Mo")
                        time.sleep(0.03)
                        send("re--")
                        if channel.recv(1) != b" ":
                            raise AssertionError("Pager was not advanced using space")
                        send("\b" * 8 + " " * 8 + "\b" * 8 + "GigabitEthernet0/1 up\r\n")
                    elif command == "display big":
                        send("Port running normally\r\n" * 30000)
                    elif command == "display slow":
                        send("first chunk\r\n")
                        time.sleep(1.6)
                        send("last chunk after quiet gap\r\n")
                    elif command == "display stall":
                        send("partial output, no completion prompt\r\n")
                        time.sleep(2.5)
                        continue
                    elif command == "display disconnect":
                        send("partial before disconnect\r\n")
                        channel.close()
                        return
                    elif command == "display invalid":
                        channel.sendall(b"before\x00padding\xffafter\r\n")
                    elif command == "display stderr":
                        channel.sendall_stderr(b"device diagnostic on stderr\r\n")
                    else:
                        send("ok\r\n")
                    send(self.prompt)
        except (EOFError, OSError, paramiko.SSHException):
            pass
        except Exception as exc:
            self.errors.append(repr(exc))
        finally:
            transport.close()

    def close(self):
        self.stop.set()
        self.sock.close()
        for connection in self.connections:
            connection.close()
        self.thread.join(2)
        for thread in self.threads:
            thread.join(4)


def workbook(path, servers, commands=COMMANDS, passwords=None, timeout=8, shared=False, legacy=False, protocols=None, ports=None):
    w = Workbook()
    w.active.title = "设备清单"
    headers = ["ip地址", "用户名", "密码", "OS类型", "端口", "编码", "超时秒", "登录方式"]
    if legacy:
        headers = headers[:4] if legacy == "four" else headers[:7]
    w.active.append(headers)
    for index, server in enumerate(servers):
        telnet = getattr(server, "protocol", "SSH") == "TELNET"
        row = [server.address, getattr(server, "username", "admin"), passwords[index] if passwords else getattr(server, "password", " secret "), "测试OS", ports[index] if ports else server.port, server.encoding if telnet or server.encoding == "big5" else "auto", timeout, protocols[index] if protocols else "TELNET" if telnet else "SSH"]
        w.active.append(row[:len(headers)])
    sheet = w.create_sheet("测试OS")
    sheet.append(["说明", "命令"])
    for command in commands:
        sheet.append(["", command])
    w.save(path)
    if shared:
        # Re-encode inline strings as shared strings, with namespaces, rich text and entities.
        from xml.etree import ElementTree as ET
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        strings = []
        with zipfile.ZipFile(path) as z:
            parts = {name: z.read(name) for name in z.namelist()}
        for name in parts:
            if not name.startswith("xl/worksheets/"):
                continue
            root = ET.fromstring(parts[name])
            for cell in root.iter(f"{{{ns}}}c"):
                if cell.get("t") == "inlineStr":
                    text = "".join(t.text or "" for t in cell.iter(f"{{{ns}}}t"))
                    for child in list(cell):
                        cell.remove(child)
                    cell.set("t", "s")
                    ET.SubElement(cell, f"{{{ns}}}v").text = str(len(strings))
                    strings.append(text)
            parts[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        root = ET.Element(f"{{{ns}}}sst")
        for text in strings:
            si = ET.SubElement(root, f"{{{ns}}}si")
            for piece in (text[:2], text[2:]):
                run = ET.SubElement(si, f"{{{ns}}}r")
                ET.SubElement(run, f"{{{ns}}}t").text = piece
        parts["xl/sharedStrings.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in parts.items():
                z.writestr(name, data)


class Integration(unittest.TestCase):
    def setUp(self):
        self.servers = []
        self.directory = ARTIFACTS / self._testMethodName
        self.directory.mkdir(parents=True, exist_ok=True)

    def tearDown(self):
        for server in self.servers:
            server.close()
        for server in self.servers:
            self.assertFalse(server.errors, server.errors)

    def server(self, *args, **kwargs):
        server = DeviceServer(*args, **kwargs)
        self.servers.append(server)
        return server

    def telnet(self, *args, **kwargs):
        server = TelnetDeviceServer(*args, **kwargs)
        self.servers.append(server)
        return server

    def run_collect(self, commands=COMMANDS, expected=0, workers=2, **kwargs):
        path = self.directory / "输入模板.xlsx"
        workbook(path, self.servers, commands, **kwargs)
        log_directory = self.directory / "输出结果"
        old_subdirectory_logs = set((log_directory / "logs").glob("collection_*.log"))
        had_subdirectory = (log_directory / "logs").exists()
        previous_logs = set(log_directory.glob("collection_*.log"))
        app_directory = ROOT / "build" / "logs"
        previous_apps = set(app_directory.glob("application_*.log"))
        start = time.perf_counter()
        result = subprocess.run([str(EXE), "--input", str(path), "--output", str(self.directory / "输出结果"), "--workers", str(workers)], capture_output=True, encoding="utf-8", timeout=40)
        elapsed = time.perf_counter() - start
        (self.directory / "collector.log").write_text(result.stdout + result.stderr, encoding="utf-8")
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        new_logs = set(log_directory.glob("collection_*.log")) - previous_logs
        self.assertEqual(len(new_logs), 1)
        self.latest_log = new_logs.pop()
        self.assertEqual(self.latest_log.parent, log_directory)
        self.assertEqual(set((log_directory / "logs").glob("collection_*.log")), old_subdirectory_logs)
        if not had_subdirectory:
            self.assertFalse((log_directory / "logs").exists())
        text = self.latest_log.read_text(encoding="utf-8-sig")
        self.assertTrue(self.latest_log.read_bytes().startswith(b"\xef\xbb\xbf"))
        self.assertIn("RUN_START", text)
        self.assertIn(f"RUN_END code={expected}", text)
        self.assertRegex(text, r"\[\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}\.\d{3}\] \[INFO ")
        self.assertRegex(text, r"pid=\d+ tid=\d+")
        app_logs = set(app_directory.glob("application_*.log")) - previous_apps
        self.assertEqual(len(app_logs), 1)
        app_text = app_logs.pop().read_text(encoding="utf-8-sig")
        self.assertIn("APPLICATION_START", app_text)
        self.assertIn(f"APPLICATION_EXIT code={expected}", app_text)
        for password in kwargs.get("passwords") or [getattr(server, "password", " secret ") for server in self.servers]:
            self.assertNotIn(password, text)
            self.assertNotIn(password, app_text)
        return elapsed, result

    def output(self, server, command):
        address = server.address.replace(":", "_")
        return self.directory / "输出结果" / f"{server.name}_{address}_{server.port}" / (command.replace("|", "_") + ".txt")

    def test_parallel_order_utf8_gbk_big5_keyboard_and_paging(self):
        a = self.server("SW-华讯", encoding="gbk")
        b = self.server("Router-2", style="hash", keyboard=True)
        c = self.server("H3C-3", style="square", encoding="big5")
        elapsed, _ = self.run_collect()
        for server in (a, b, c):
            self.assertEqual(server.commands, COMMANDS)
            version = self.output(server, COMMANDS[0]).read_text(encoding="utf-8-sig")
            self.assertIn("接口狀態正常" if server.encoding == "big5" else "接口状态正常", version)
            self.assertNotIn("\ufffd", version)
            self.assertNotIn("\x1b", version)
            self.assertIn("GigabitEthernet0/1 up", self.output(server, COMMANDS[1]).read_text(encoding="utf-8-sig"))
            self.assertEqual(self.output(server, COMMANDS[2]).read_text(encoding="utf-8-sig").count("Port running normally"), 30000)
            self.assertIn("last chunk after quiet gap", self.output(server, COMMANDS[3]).read_text(encoding="utf-8-sig"))
        self.assertLess(abs(a.timestamps[0] - b.timestamps[0]), 1.5, "First two devices did not overlap")
        self.assertGreater(c.timestamps[0], min(a.timestamps[-1], b.timestamps[-1]), "Worker bound was not respected")
        with (self.directory / "输出结果" / "_summary.csv").open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual([row["状态"] for row in rows], ["SUCCESS"] * 3)
        (self.directory / "metrics.json").write_text(json.dumps({"devices": 3, "workers": 2, "elapsed_seconds": round(elapsed, 3), "commands": 12, "large_output_lines": 90000}), encoding="utf-8")

    def test_shared_strings_and_filename_collisions(self):
        a = self.server("CON", style="hash")
        commands = ["show a | b", "show a _ b", "show a | b", "show 中文 & status", "show État", "SHOW état", "display stderr"]
        self.run_collect(commands=commands, shared=True)
        self.assertEqual(a.commands, commands)
        folder = self.output(a, commands[0]).parent
        self.assertTrue((folder / "show a _ b (2).txt").exists())
        self.assertTrue((folder / "show a _ b (3).txt").exists())
        self.assertTrue((folder / "show 中文 & status.txt").exists())
        self.assertTrue((folder / "SHOW état (2).txt").exists())
        self.assertIn("device diagnostic on stderr", (folder / "display stderr.txt").read_text(encoding="utf-8-sig"))

    def test_timeout_preserves_partial_and_stops_sequence(self):
        a = self.server("Timeout")
        self.run_collect(commands=["display stall", "must not run"], timeout=1, expected=1)
        self.assertEqual(a.commands, ["display stall"])
        self.assertIn("INCOMPLETE", self.output(a, "display stall").read_text(encoding="utf-8-sig"))

    def test_disconnect_is_failure(self):
        a = self.server("Disconnect")
        self.run_collect(commands=["display disconnect", "must not run"], expected=1)
        self.assertEqual(a.commands, ["display disconnect"])
        self.assertIn("INCOMPLETE", self.output(a, "display disconnect").read_text(encoding="utf-8-sig"))

    def test_wrong_password(self):
        a = self.server("AuthFail")
        self.run_collect(passwords=["wrong-password"], expected=1)
        self.assertEqual(a.commands, [])
        log = self.latest_log.read_text(encoding="utf-8-sig")
        self.assertIn("SSH authentication failed", log)
        self.assertIn("[ERROR]", log)

    def test_host_key_change_rejected(self):
        a = self.server("Pinned")
        self.run_collect(commands=["display version"])
        a.commands.clear()
        a.key = paramiko.RSAKey.generate(2048)
        _, result = self.run_collect(commands=["display version"], expected=1)
        self.assertIn("host key changed", result.stdout)
        self.assertEqual(a.commands, [])

    def test_ecdsa_and_ipv6(self):
        a = self.server("IPv6-Core", style="hash", key=paramiko.ECDSAKey.generate(bits=256), address="::1")
        self.run_collect(commands=["display version"])
        self.assertIn("接口状态正常", self.output(a, "display version").read_text(encoding="utf-8-sig"))

    def test_long_unicode_paths(self):
        a = self.server("LongPath")
        self.directory = Path("\\\\?\\" + str(self.directory / ("长目录" * 20) / ("nested" * 15)))
        self.directory.mkdir(parents=True, exist_ok=True)
        self.run_collect(commands=["display version"])
        self.assertIn("接口状态正常", self.output(a, "display version").read_text(encoding="utf-8-sig"))

    def test_invalid_bytes_preserve_original(self):
        a = self.server("RawOutput")
        _, result = self.run_collect(commands=["display invalid"])
        output = self.output(a, "display invalid")
        self.assertIn("after", output.read_text(encoding="utf-8-sig"))
        self.assertIn(b"\x00padding\xffafter", output.with_suffix(".txt.raw").read_bytes())
        self.assertIn("编码转换有损", result.stdout)

    def test_runtime_logs_redact_password_and_keep_runs(self):
        a = self.server("Log-Switch")
        commands = ["show include secret token"]
        self.run_collect(commands=commands)
        first = self.latest_log
        self.assertIn("[REDACTED]", first.read_text(encoding="utf-8-sig"))
        session = self.output(a, commands[0]).parent / "_session.log"
        self.assertNotIn(" secret ", session.read_text(encoding="utf-8-sig"))
        self.run_collect(commands=["display version"])
        self.assertNotEqual(first, self.latest_log)
        self.assertTrue(first.exists())

    def test_runtime_log_is_visible_before_task_finishes(self):
        self.server("LiveLog")
        path = self.directory / "live.xlsx"
        workbook(path, self.servers, ["display stall"], timeout=2)
        output = self.directory / "live-output"
        previous_logs = set(output.glob("collection_*.log"))
        process = subprocess.Popen([str(EXE), "--input", str(path), "--output", str(output)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
        try:
            deadline = time.monotonic() + 8
            observed = False
            while time.monotonic() < deadline and process.poll() is None:
                for log in set(output.glob("collection_*.log")) - previous_logs:
                    text = log.read_text(encoding="utf-8-sig")
                    if "COMMAND_START" in text:
                        self.assertNotIn("RUN_END", text)
                        observed = True
                        break
                if observed:
                    break
                time.sleep(0.02)
            self.assertTrue(observed, "Runtime log was not readable during collection")
            process.communicate(timeout=8)
            self.assertEqual(process.returncode, 1)
        finally:
            if process.poll() is None:
                process.kill()
                process.communicate(timeout=5)

    def test_template_failure_is_logged_and_no_log_directory_is_error(self):
        path = self.directory / "invalid.xlsx"
        w = Workbook()
        w.active.append(["ip地址"])
        w.save(path)
        output = self.directory / "template-error"
        result = subprocess.run([str(EXE), "--input", str(path), "--output", str(output)], capture_output=True, timeout=8)
        self.assertEqual(result.returncode, 2)
        log = max(output.glob("collection_*.log"), key=lambda item: item.stat().st_mtime_ns).read_text(encoding="utf-8-sig")
        self.assertIn("TEMPLATE_FAILED", log)
        self.assertIn("RUN_END code=2", log)
        blocked = self.directory / "blocked-output-file"
        blocked.write_text("not a directory", encoding="utf-8")
        result = subprocess.run([str(EXE), "--input", str(path), "--output", str(blocked)], capture_output=True, encoding="utf-8", timeout=8)
        self.assertEqual(result.returncode, 2)
        self.assertIn("无法创建输出目录", result.stdout)

    def test_telnet_negotiation_and_mixed_parallel_collection(self):
        ssh = self.server("SSH-Mixed")
        telnet = self.telnet("Telnet-核心", encoding="gbk", mode="chinese", echo_password=True)
        password_only = self.telnet("Telnet-Password", mode="password", username="")
        self.run_collect()
        for server in self.servers:
            self.assertEqual(server.commands, COMMANDS)
            version = self.output(server, COMMANDS[0]).read_text(encoding="utf-8-sig")
            self.assertIn("接口状态正常", version)
            self.assertNotIn(" secret ", version)
            self.assertNotIn("\ufffd", version)
            self.assertIn("GigabitEthernet0/1 up", self.output(server, COMMANDS[1]).read_text(encoding="utf-8-sig"))
            self.assertEqual(self.output(server, COMMANDS[2]).read_text(encoding="utf-8-sig").count("Port running normally"), 30000)
            self.assertIn("last chunk after quiet gap", self.output(server, COMMANDS[3]).read_text(encoding="utf-8-sig"))
        self.assertLess(abs(ssh.timestamps[0] - telnet.timestamps[0]), 1.5)
        for server in (telnet, password_only):
            self.assertIn((254, 99), server.negotiations)  # DONT unsupported WILL.
            self.assertIn((252, 100), server.negotiations)  # WONT unsupported DO.
            self.assertIn((24, b"\x00VT100"), server.subnegotiations)
            self.assertIn((31, b"\x00\xf0\x03\xe8"), server.subnegotiations)
        with (self.directory / "输出结果" / "_summary.csv").open(encoding="utf-8-sig", newline="") as f:
            self.assertEqual([row["登录方式"] for row in csv.DictReader(f)], ["SSH", "TELNET", "TELNET"])
        self.assertIn("TELNET_AUTH_OK", self.latest_log.read_text(encoding="utf-8-sig"))

    def test_telnet_login_variants_and_iac_escaping(self):
        login = self.telnet("Telnet-Login", mode="login")
        chinese = self.telnet("Telnet-GBK", encoding="gbk", mode="chinese", username="运维员", password="口令123", echo_password=True)
        latin = self.telnet("Telnet-Latin", encoding="latin1", password="pÿss", echo_password=True)
        direct = self.telnet("Telnet-Direct", mode="direct")
        self.run_collect(commands=["show iac"])
        for server in (login, chinese, latin, direct):
            self.assertEqual(server.commands, ["show iac"])
        self.assertIn("literal=ÿ", self.output(latin, "show iac").read_text(encoding="utf-8-sig"))
        self.assertIn(b"p\xff\xffss\r\n", latin.received_wire)
        self.assertIn("运维员".encode("gbk"), chinese.received_wire)
        self.assertNotIn("口令123", self.latest_log.read_text(encoding="utf-8-sig"))

    def test_telnet_failed_authentication_repeats_and_password_change(self):
        rejected = self.telnet("Telnet-Reject", mode="reject")
        repeated = self.telnet("Telnet-Repeat", mode="repeat")
        change = self.telnet("Telnet-Change", mode="change")
        self.run_collect(commands=["must not run"], expected=1, workers=3)
        for server in (rejected, repeated, change):
            self.assertEqual(server.commands, [])
        log = self.latest_log.read_text(encoding="utf-8-sig")
        self.assertIn("Telnet authentication failed", log)
        self.assertIn("repeated credential prompt", log)
        self.assertIn("password change is required", log)

    def test_telnet_login_timeout_and_missing_username(self):
        timeout = self.telnet("Telnet-LoginTimeout", mode="timeout")
        missing = self.telnet("Telnet-NoUsername", username="")
        self.run_collect(commands=["must not run"], timeout=1, expected=1)
        self.assertEqual(timeout.commands, [])
        self.assertEqual(missing.commands, [])
        log = self.latest_log.read_text(encoding="utf-8-sig")
        self.assertIn("Timed out waiting", log)
        self.assertIn("requested a username", log)

    def test_telnet_disconnect_and_timeout_preserve_partial(self):
        disconnect = self.telnet("Telnet-Disconnect", close_on_command=True)
        timeout = self.telnet("Telnet-CmdTimeout")
        commands = ["display stall", "must not run"]
        self.run_collect(commands=commands, timeout=1, expected=1)
        for server in (disconnect, timeout):
            self.assertEqual(server.commands, ["display stall"])
            self.assertIn("INCOMPLETE", self.output(server, commands[0]).read_text(encoding="utf-8-sig"))

    def test_protocol_defaults_legacy_and_invalid_values(self):
        server = self.server("InventoryOnly")
        path = self.directory / "protocol.xlsx"
        inspector = ROOT / "build" / "inspect_inventory.exe"
        for protocol, legacy, expected in (("SSH", False, 22), ("TELNET", False, 23), ("telnet", False, 23), (None, False, 22), ("TELNET", True, 22), ("TELNET", "four", 22)):
            workbook(path, [server], protocols=[protocol], ports=[None], legacy=legacy)
            result = subprocess.run([str(inspector), str(path)], capture_output=True, encoding="utf-8", timeout=5)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(json.loads(result.stdout), [{"protocol": "SSH" if expected == 22 else "TELNET", "port": expected}])
        workbook(path, [server], protocols=["TELNET"], ports=[2323])
        result = subprocess.run([str(inspector), str(path)], capture_output=True, encoding="utf-8", timeout=5)
        self.assertEqual(json.loads(result.stdout)[0]["port"], 2323)
        workbook(path, [server], protocols=["RLOGIN"])
        result = subprocess.run([str(EXE), "--check", str(path)], capture_output=True, encoding="utf-8", timeout=5)
        self.assertEqual(result.returncode, 2)
        self.assertIn("SSH/TELNET", result.stderr)
        workbook(path, [server], protocols=["TELNET"], passwords=["password\ninjected"])
        result = subprocess.run([str(inspector), str(path)], capture_output=True, timeout=5)
        self.assertEqual(result.returncode, 2)

    def test_telnet_gui_success_and_cancel(self):
        if not RUN_GUI:
            self.skipTest("GUI tests require --gui")
        a = self.telnet("Telnet-GUI")
        b = self.server("SSH-Queued")
        path = self.directory / "Telnet界面测试.xlsx"
        for mode, commands in (("success", ["display version"]), ("cancel", ["display stall", "must not run"])):
            a.commands.clear()
            workbook(path, [a, b] if mode == "cancel" else [a], commands)
            output = self.directory / mode
            screenshot = self.directory / f"gui-{mode}.bmp"
            result = subprocess.run([str(ROOT / "build" / "gui_tests.exe"), str(path), str(output), str(screenshot), mode], capture_output=True, encoding="utf-8", timeout=35)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(a.commands, commands if mode == "success" else ["display stall"])
            self.assertEqual(b.commands, [])
            with (output / "_summary.csv").open(encoding="utf-8-sig", newline="") as f:
                self.assertEqual([row["状态"] for row in csv.DictReader(f)], ["SUCCESS"] if mode == "success" else ["CANCELLED", "CANCELLED"])

    def test_invalid_templates(self):
        a = self.server("Unused")
        path = self.directory / "invalid.xlsx"
        w = Workbook()
        w.active.append(["ip地址"])
        w.save(path)
        result = subprocess.run([str(EXE), "--check", str(path)], capture_output=True, encoding="utf-8", timeout=5)
        self.assertEqual(result.returncode, 2)
        workbook(path, [a, a])
        result = subprocess.run([str(EXE), "--check", str(path)], capture_output=True, encoding="utf-8", timeout=5)
        self.assertEqual(result.returncode, 2)
        self.assertIn("Duplicate", result.stderr)

    def test_gui_success_and_cancel(self):
        if not RUN_GUI:
            self.skipTest("GUI tests require --gui")
        a = self.server("GUI-Switch")
        b = self.server("GUI-Queued")
        path = self.directory / "图形界面测试.xlsx"
        for mode, commands in (("success", ["display version"]), ("cancel", ["display stall", "must not run"])):
            a.commands.clear()
            workbook(path, [a, b] if mode == "cancel" else [a], commands, timeout=8)
            output = self.directory / mode
            image = self.directory / f"gui-{mode}.bmp"
            result = subprocess.run([str(ROOT / "build" / "gui_tests.exe"), str(path), str(output), str(image), mode], capture_output=True, timeout=35, encoding="utf-8")
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            if mode == "success":
                self.assertEqual(a.commands, commands)
            else:
                self.assertEqual(a.commands, ["display stall"])
                self.assertEqual(b.commands, [])
                with (output / "_summary.csv").open(encoding="utf-8-sig", newline="") as f:
                    self.assertEqual([row["状态"] for row in csv.DictReader(f)], ["CANCELLED", "CANCELLED"])
            try:
                from PIL import Image
                Image.open(image).save(image.with_suffix(".png"))
            except ImportError:
                pass


if __name__ == "__main__":
    import sys
    if "--gui" in sys.argv:
        RUN_GUI = True
        sys.argv.remove("--gui")
    ARTIFACTS.mkdir(exist_ok=True)
    unittest.main(verbosity=2)

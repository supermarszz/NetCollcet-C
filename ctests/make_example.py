"""Generate a credential-free example for the native C release package."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


def main():
    root = Path(__file__).resolve().parents[1]
    workbook = Workbook()
    devices = workbook.active
    devices.title = "设备清单"
    devices.append(["ip地址", "用户名", "密码", "OS类型", "端口", "编码", "超时秒", "登录方式"])
    devices.append(["192.0.2.10", "admin", "请填写实际密码", "Huawei", 22, "auto", 90, "SSH"])
    devices.append(["192.0.2.11", "admin", "请填写实际密码", "Cisco", 22, "utf-8", 90, "SSH"])
    devices.append(["192.0.2.12", "admin", "请填写实际密码", "H3C", 23, "gbk", 90, "TELNET"])
    examples = {
        "Huawei": ["display version", "display device", "display interface brief", "display current-configuration"],
        "H3C": ["display version", "display device", "display interface brief"],
        "Cisco": ["show version", "show inventory", "show interfaces", "show running-config"],
        "Juniper": ["show version", "show chassis hardware", "show interfaces terse"],
    }
    for name, commands in examples.items():
        sheet = workbook.create_sheet(name)
        sheet.append(["命令"])
        for command in commands:
            sheet.append([command])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="245C45")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 26
        for column in sheet.columns:
            for cell in column:
                cell.number_format = "@" if cell.column in (1, 2, 3, 4, 6, 8) else "General"
            sheet.column_dimensions[column[0].column_letter].width = 25 if sheet is devices else 55
    validation = DataValidation(type="list", formula1='"auto,utf-8,gb18030,gbk,big5,latin1"')
    devices.add_data_validation(validation)
    validation.add("F2:F10000")
    protocol = DataValidation(type="list", formula1='"SSH,TELNET"', allow_blank=True)
    protocol.errorTitle = "登录方式错误"
    protocol.error = "请选择 SSH 或 TELNET；留空默认 SSH。"
    protocol.showErrorMessage = True
    protocol.errorStyle = "stop"
    devices.add_data_validation(protocol)
    protocol.add("H2:H10000")
    workbook.save(root / "C版采集模板.xlsx")


if __name__ == "__main__":
    main()
